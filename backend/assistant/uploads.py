"""Upload text extraction for the assistant.

Extracts readable text from uploaded PDF / DOCX / XLSX / CSV / plain-text files
so the assistant can read them. Extracted text is wrapped in untrusted-content
delimiters (assistant/delimiters.py) and prepended into the user's message by the
router — there is no attachments table and no file cache (nothing forwards the
original bytes anywhere). Adapted from Chatty's ``core/agents/tools/text_extraction.py``.

Heavy parsers (pdfplumber / python-docx / openpyxl) are imported LAZILY inside
the functions that use them, so the assistant package imports cleanly even when
those libraries are absent (CI import check, minimal installs).
"""

import io

from assistant import delimiters

ALLOWED_EXTENSIONS = {"csv", "xlsx", "md", "txt", "pdf", "docx"}
MAX_FILES = 5
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_CHARS = 50_000  # per-file extracted-text cap


class UploadError(ValueError):
    """A user-facing upload problem (bad extension, unreadable/corrupt file).

    The router maps this to HTTP 400 with the message shown to the user.
    """


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""


def extract_pdf_text(data: bytes, max_chars: int) -> str:
    import pdfplumber  # lazy: optional dependency

    parts: list[str] = []
    total = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if not page_text:
                    continue
                remaining = max_chars - total
                if remaining <= 0:
                    break
                parts.append(page_text[:remaining])
                total += min(len(page_text), remaining)
    except Exception as e:  # corrupt / not a real PDF
        raise UploadError(f"Could not read PDF: {e}") from e
    return "\n\n".join(parts)


def extract_docx_text(data: bytes, max_chars: int) -> str:
    import docx  # lazy: optional dependency (python-docx)

    parts: list[str] = []
    total = 0
    try:
        doc = docx.Document(io.BytesIO(data))
        rows = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                rows.append("\t".join(cell.text.strip() for cell in row.cells))
        for text in rows:
            if not text:
                continue
            remaining = max_chars - total
            if remaining <= 0:
                break
            parts.append(text[:remaining])
            total += min(len(text), remaining) + 1
    except Exception as e:
        raise UploadError(f"Could not read Word document: {e}") from e
    return "\n".join(parts)


def extract_xlsx_text(data: bytes, max_chars: int) -> str:
    from openpyxl import load_workbook  # lazy: optional dependency

    parts: list[str] = []
    total = 0
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if len(wb.sheetnames) > 1:
                    header = f"--- Sheet: {sheet_name} ---"
                    parts.append(header)
                    total += len(header)
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join("" if c is None else str(c) for c in row)
                    remaining = max_chars - total
                    if remaining <= 0:
                        break
                    parts.append(line[:remaining])
                    total += min(len(line), remaining) + 1
                if total >= max_chars:
                    break
        finally:
            wb.close()
    except UploadError:
        raise
    except Exception as e:
        raise UploadError(f"Could not read spreadsheet: {e}") from e
    return "\n".join(parts)


def extract_plain(data: bytes, max_chars: int) -> str:
    # latin-1 decodes any byte sequence, so encoding never fails; guard against
    # binary content (which would decode to garbage) by rejecting NUL bytes.
    if b"\x00" in data:
        raise UploadError("File appears to be binary, not text.")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return text[:max_chars]


def _extract_by_ext(ext: str, data: bytes, max_chars: int) -> str:
    if ext == "pdf":
        return extract_pdf_text(data, max_chars)
    if ext == "docx":
        return extract_docx_text(data, max_chars)
    if ext == "xlsx":
        return extract_xlsx_text(data, max_chars)
    # csv / md / txt
    return extract_plain(data, max_chars)


def extract_upload(filename: str, data: bytes) -> str:
    """Extract text from one uploaded file and return a wrapped, tagged block.

    Raises UploadError (→ HTTP 400) on a disallowed extension or an unreadable
    file. Truncated extractions get a footer; empty extractions get a placeholder
    (still wrapped, so the model sees the filename).
    """
    ext = _ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise UploadError(f"Unsupported file type '.{ext or filename}'. Allowed: {allowed}.")

    text = _extract_by_ext(ext, data, MAX_CHARS)

    if len(text) >= MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n(truncated at 50,000 characters)"
    elif not text.strip():
        text = "(No extractable text found.)"

    return delimiters.wrap_untrusted_file(filename, text)
